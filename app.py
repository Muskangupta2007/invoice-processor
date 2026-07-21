import os
import io
from flask import Flask, request, jsonify, send_from_directory, send_file
from werkzeug.utils import secure_filename
from datetime import datetime
import subprocess
import pandas as pd
from functools import wraps
import db
import extractor
import teams

import gzip

app = Flask(__name__, static_folder='static', static_url_path='')

UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024

db.init_db()

# --- AUTHENTICATION DECORATOR ---
def requires_auth(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        # Read from Authorization header or from request query args (for standard link downloads)
        token = request.headers.get("Authorization")
        if token and token.startswith("Bearer "):
            token = token.split(" ")[1]
        else:
            token = request.args.get("token")
            
        expected_token = db.get_setting("admin_password", "invosync2026")
        if not token or token != expected_token:
            return jsonify({"status": "error", "message": "Unauthorized access. Please login."}), 401
        return f(*args, **kwargs)
    return decorated

def requires_settings_auth(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        token = request.headers.get("Authorization")
        if token and token.startswith("Bearer "):
            token = token.split(" ")[1]
        else:
            token = request.args.get("token")
            
        expected_token = db.get_setting("settings_password", "settings2026")
        if not token or token != expected_token:
            return jsonify({"status": "error", "message": "Incorrect Settings Lock password."}), 401
        return f(*args, **kwargs)
    return decorated

@app.route('/')
def index():
    return send_from_directory('static', 'index.html')

@app.after_request
def add_header(response):
    path = request.path.lower()
    # 1. Cache static media assets (images/fonts) for 24 hours, but revalidate JS/CSS
    if any(path.endswith(ext) for ext in ['.png', '.jpg', '.jpeg', '.svg', '.ico', '.woff', '.woff2']):
        response.headers["Cache-Control"] = "public, max-age=86400"
    else:
        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        response.headers["Pragma"] = "no-cache"

    # 2. Gzip compression for API JSON and static responses
    accept_encoding = request.headers.get('Accept-Encoding', '')
    if (
        'gzip' in accept_encoding.lower()
        and response.status_code == 200
        and not response.direct_passthrough
    ):
        content_type = response.content_type or ''
        if any(ct in content_type for ct in ['json', 'html', 'css', 'javascript', 'text']):
            data = response.get_data()
            if len(data) > 500:
                compressed_data = gzip.compress(data)
                response.set_data(compressed_data)
                response.headers['Content-Encoding'] = 'gzip'
                response.headers['Content-Length'] = len(compressed_data)
                response.headers['Vary'] = 'Accept-Encoding'

    return response

# --- AUTH ENDPOINTS ---

@app.route('/api/login', methods=['POST'])
def login():
    data = request.json or {}
    username = data.get("username")
    password = data.get("password")
    
    expected_username = db.get_setting("admin_username", "admin")
    expected_password = db.get_setting("admin_password", "invosync2026")
    
    if username == expected_username and password == expected_password:
        return jsonify({
            "status": "success",
            "message": "Login successful.",
            "token": expected_password
        })
    else:
        return jsonify({"status": "error", "message": "Invalid username or password."}), 401

# --- SETTINGS ENDPOINTS ---

@app.route('/api/settings/status', methods=['GET'])
@requires_auth
def get_settings_status():
    api_key = db.get_setting("azure_api_key", "")
    webhook_url = db.get_setting("teams_webhook_url", "")
    return jsonify({
        "status": "success",
        "azure_api_key_configured": bool(api_key),
        "teams_webhook_url": bool(webhook_url)
    })

@app.route('/api/settings', methods=['GET'])
@requires_settings_auth
def get_settings():
    azure_endpoint = db.get_setting("azure_endpoint", "")
    api_key = db.get_setting("azure_api_key", "")
    webhook_url = db.get_setting("teams_webhook_url", "")
    excel_dir = db.get_setting("excel_storage_directory", "")
    admin_user = db.get_setting("admin_username", "admin")
    
    # Google Sheets settings
    gsheet_url_sk_agro = db.get_setting("gsheet_url_sk_agro", "")
    gsheet_sync_url_sk_agro = db.get_setting("gsheet_sync_url_sk_agro", "")
    gsheet_url_sk_electricals = db.get_setting("gsheet_url_sk_electricals", "")
    gsheet_sync_url_sk_electricals = db.get_setting("gsheet_sync_url_sk_electricals", "")
    gsheet_url_triotech = db.get_setting("gsheet_url_triotech", "")
    gsheet_sync_url_triotech = db.get_setting("gsheet_sync_url_triotech", "")
    
    masked_key = ""
    if api_key:
        masked_key = api_key[:4] + "*" * (len(api_key) - 8) + api_key[-4:] if len(api_key) > 8 else "****"
        
    return jsonify({
        "azure_endpoint": azure_endpoint,
        "azure_api_key_configured": bool(api_key),
        "azure_api_key_masked": masked_key,
        "teams_webhook_url": webhook_url,
        "excel_storage_directory": excel_dir,
        "admin_username": admin_user,
        
        # Include Google Sheets settings in response
        "gsheet_url_sk_agro": gsheet_url_sk_agro,
        "gsheet_sync_url_sk_agro": gsheet_sync_url_sk_agro,
        "gsheet_url_sk_electricals": gsheet_url_sk_electricals,
        "gsheet_sync_url_sk_electricals": gsheet_sync_url_sk_electricals,
        "gsheet_url_triotech": gsheet_url_triotech,
        "gsheet_sync_url_triotech": gsheet_sync_url_triotech
    })

@app.route('/api/settings', methods=['POST'])
@requires_settings_auth
def save_settings():
    data = request.json or {}
    azure_endpoint = data.get("azure_endpoint")
    api_key = data.get("azure_api_key")
    webhook_url = data.get("teams_webhook_url")
    excel_dir = data.get("excel_storage_directory")
    
    # Google Sheets settings from payload
    gsheet_url_sk_agro = data.get("gsheet_url_sk_agro")
    gsheet_sync_url_sk_agro = data.get("gsheet_sync_url_sk_agro")
    gsheet_url_sk_electricals = data.get("gsheet_url_sk_electricals")
    gsheet_sync_url_sk_electricals = data.get("gsheet_sync_url_sk_electricals")
    gsheet_url_triotech = data.get("gsheet_url_triotech")
    gsheet_sync_url_triotech = data.get("gsheet_sync_url_triotech")
    
    admin_user = data.get("admin_username")
    admin_pass = data.get("admin_password")
    settings_lock_pass = data.get("settings_lock_password")
    current_pass = data.get("current_password")
    
    if admin_user or admin_pass or settings_lock_pass:
        expected_pass = db.get_setting("admin_password", "invosync2026")
        if not current_pass or current_pass != expected_pass:
            return jsonify({"status": "error", "message": "Verification failed: Incorrect current password. Settings not saved."}), 400
            
        if admin_user:
            db.save_setting("admin_username", admin_user.strip())
        if admin_pass:
            db.save_setting("admin_password", admin_pass.strip())
        if settings_lock_pass:
            db.save_setting("settings_password", settings_lock_pass.strip())
            
    if azure_endpoint is not None:
        db.save_setting("azure_endpoint", azure_endpoint.strip())
        
    if api_key and "*" in api_key:
        pass
    elif api_key is not None:
        db.save_setting("azure_api_key", api_key.strip())
        
    if webhook_url is not None:
        db.save_setting("teams_webhook_url", webhook_url.strip())
        
    if excel_dir is not None:
        db.save_setting("excel_storage_directory", excel_dir.strip())
        
    # Save Google Sheets settings
    if gsheet_url_sk_agro is not None:
        db.save_setting("gsheet_url_sk_agro", gsheet_url_sk_agro.strip())
    if gsheet_sync_url_sk_agro is not None:
        db.save_setting("gsheet_sync_url_sk_agro", gsheet_sync_url_sk_agro.strip())
        
    if gsheet_url_sk_electricals is not None:
        db.save_setting("gsheet_url_sk_electricals", gsheet_url_sk_electricals.strip())
    if gsheet_sync_url_sk_electricals is not None:
        db.save_setting("gsheet_sync_url_sk_electricals", gsheet_sync_url_sk_electricals.strip())
        
    if gsheet_url_triotech is not None:
        db.save_setting("gsheet_url_triotech", gsheet_url_triotech.strip())
    if gsheet_sync_url_triotech is not None:
        db.save_setting("gsheet_sync_url_triotech", gsheet_sync_url_triotech.strip())
        
    try:
        db.sync_to_excel()
    except Exception as e:
        print(f"Error syncing excel files on settings change: {e}")
        
    # Trigger initial push sync to Google Sheets if sync URLs were updated
    import gsheet
    import threading
    for comp in db.COMPANIES:
        comp_key = comp.lower().replace(' ', '_')
        if data.get(f"gsheet_sync_url_{comp_key}") is not None and data.get(f"gsheet_sync_url_{comp_key}").strip() != "":
            try:
                threading.Thread(target=gsheet.sync_to_google_sheet, args=(comp,), daemon=True).start()
            except Exception as e:
                print(f"Error starting initial Google Sheets sync for {comp}: {e}")
                
    return jsonify({"status": "success", "message": "Settings updated successfully."})

@app.route('/api/settings/test-teams', methods=['POST'])
@requires_settings_auth
def test_teams():
    webhook_url = db.get_setting("teams_webhook_url")
    if not webhook_url:
        return jsonify({"status": "error", "message": "Microsoft Teams Webhook URL is not configured."}), 400
        
    success = teams.send_teams_notification(webhook_url, "test", {})
    if success:
        return jsonify({"status": "success", "message": "Test notification sent to Microsoft Teams."})
    else:
        return jsonify({"status": "error", "message": "Failed to send test notification. Check server logs."}), 500

# --- LEDGER ENDPOINTS ---

@app.route('/api/invoices', methods=['GET'])
@requires_auth
def get_documents():
    company = request.args.get('company')
    doc_type = request.args.get('doc_type')
    
    try:
        documents = db.get_documents(company, doc_type)
        for doc in documents:
            amt = doc.get('amount') or 0.0
            if (not amt or amt == 0.0) and doc.get('doc_type') == 'Challan':
                qtys = str(doc.get('deliver_qty') or '').split('\n')
                prices = str(doc.get('unit_price') or '').split('\n')
                chal_amt = 0.0
                for q_str, p_str in zip(qtys, prices):
                    try:
                        q = float(q_str.replace(',', '').strip())
                        p = float(p_str.replace(',', '').strip())
                        chal_amt += q * p
                    except (ValueError, TypeError):
                        pass
                doc['amount'] = chal_amt
        
        sequence_check = {"status": "ok", "missing_numbers": []}
        if company and doc_type:
            sequence_check = db.check_invoice_sequence(company, doc_type)
            missing_mrns = db.check_missing_mrn(company, doc_type)
            missing_count = len(missing_mrns)
        else:
            missing_mrns = db.check_missing_mrn()
            missing_count = len(missing_mrns)
            
        return jsonify({
            "status": "success",
            "invoices": documents,
            "sequence_check": sequence_check,
            "missing_mrn_count": missing_count
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/upload', methods=['POST'])
@requires_auth
def upload_file():
    if 'file' not in request.files:
        return jsonify({"status": "error", "message": "No file part in request."}), 400
        
    file = request.files['file']
    company = request.form.get("company")
    doc_type = request.form.get("doc_type")
    
    if not company or not doc_type:
        return jsonify({"status": "error", "message": "Target Company and Document Type must be specified."}), 400
        
    if file.filename == '':
        return jsonify({"status": "error", "message": "No file selected."}), 400
        
    endpoint = db.get_setting("azure_endpoint")
    api_key = db.get_setting("azure_api_key")
    if not endpoint or not api_key:
        return jsonify({
            "status": "error", 
            "message": "Azure Document Intelligence endpoint or API key is not configured. Please configure them in settings."
        }), 400
        
    filename = secure_filename(file.filename)
    unique_filename = f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{filename}"
    file_path = os.path.join(app.config['UPLOAD_FOLDER'], unique_filename)
    file.save(file_path)
    
    try:
        extracted_data = extractor.extract_invoice_data(file_path, endpoint, api_key, doc_type)
        doc_no = extracted_data.get("doc_no")
        
        extracted_data["company"] = company
        extracted_data["doc_type"] = doc_type
        
        if not doc_no:
            return jsonify({
                "status": "error",
                "message": "Could not identify a document number from the document.",
                "extracted_data": extracted_data
            }), 422
            
        existing_doc = db.get_document_by_no(company, doc_type, doc_no)
        if existing_doc:
            return jsonify({
                "status": "duplicate",
                "message": f"{doc_type} number '{doc_no}' already exists for company '{company}'.",
                "new_data": extracted_data,
                "existing_data": existing_doc
            })
            
        success = db.add_document(extracted_data)
        if not success:
            return jsonify({"status": "error", "message": "Failed to write document to database."}), 500
            
        # Send Teams notifications
        webhook_url = db.get_setting("teams_webhook_url")
        if webhook_url:
            notify_payload = extracted_data.copy()
            notify_payload["processed_time"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            notify_payload["invoice_no"] = notify_payload["doc_no"]
            notify_payload["invoice_date"] = notify_payload["doc_date"]
            notify_payload["invoice_amt"] = notify_payload["amount"]
            
            teams.send_teams_notification(webhook_url, "new_invoice", notify_payload)
            
            seq_check = db.check_invoice_sequence(company, doc_type)
            if seq_check["status"] == "warning":
                teams.send_teams_notification(webhook_url, "sequence_gap", seq_check)
                
            if notify_payload.get("mrn_status") == "PENDING" or not notify_payload.get("mrn_no"):
                teams.send_teams_notification(webhook_url, "missing_mrn", notify_payload)
                
        return jsonify({
            "status": "success",
            "message": "Document successfully uploaded and processed.",
            "data": extracted_data
        })
        
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/invoices/confirm-upload', methods=['POST'])
@requires_auth
def confirm_upload():
    data = request.json or {}
    action = data.get("action")
    doc_data = data.get("data")
    
    if not doc_data or not action:
        return jsonify({"status": "error", "message": "Missing required fields 'action' and 'data'."}), 400
        
    company = doc_data.get("company")
    doc_type = doc_data.get("doc_type")
    doc_no = doc_data.get("doc_no")
    
    if action == "skip":
        return jsonify({"status": "skipped", "message": "Document upload skipped by user."})
        
    if action == "replace":
        success = db.update_document(company, doc_type, doc_no, doc_data)
        if not success:
            return jsonify({"status": "error", "message": "Failed to update the existing document."}), 500
            
        webhook_url = db.get_setting("teams_webhook_url")
        if webhook_url:
            notify_payload = doc_data.copy()
            notify_payload["processed_time"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S") + " (Replaced)"
            notify_payload["invoice_no"] = notify_payload["doc_no"]
            notify_payload["invoice_date"] = notify_payload["doc_date"]
            notify_payload["invoice_amt"] = notify_payload["amount"]
            
            teams.send_teams_notification(webhook_url, "new_invoice", notify_payload)
            if notify_payload.get("mrn_status") == "PENDING" or not notify_payload.get("mrn_no"):
                teams.send_teams_notification(webhook_url, "missing_mrn", notify_payload)
                
        return jsonify({"status": "success", "message": f"{doc_type} '{doc_no}' successfully replaced."})
        
    return jsonify({"status": "error", "message": f"Invalid action '{action}'."}), 400

@app.route('/api/invoices/<company>/<doc_type>/<path:doc_no>', methods=['PUT'])
@requires_auth
def update_document(company, doc_type, doc_no):
    data = request.json or {}
    if 'mrn_status' in data and data['mrn_status']:
        data['mrn_status'] = str(data['mrn_status']).strip().upper()
    try:
        existing = db.get_document_by_no(company, doc_type, doc_no)
        if not existing:
            return jsonify({"status": "error", "message": "Document not found."}), 404
            
        success = db.update_document(company, doc_type, doc_no, data)
        if success:
            old_mrn_status = existing.get("mrn_status", "PENDING")
            new_mrn_status = data.get("mrn_status", old_mrn_status)
            new_mrn_no = data.get("mrn_no", existing.get("mrn_no"))
            
            webhook_url = db.get_setting("teams_webhook_url")
            if webhook_url and (old_mrn_status != new_mrn_status or (new_mrn_status == 'DONE' and new_mrn_no != existing.get("mrn_no"))):
                updated_record = db.get_document_by_no(company, doc_type, doc_no)
                card_data = {
                    "message": f"MRN details updated for {doc_type} **{doc_no}** ({company}):\n\n* **MRN Status**: {updated_record.get('mrn_status')}\n* **MRN No**: {updated_record.get('mrn_no')}\n* **MRN Date**: {updated_record.get('mrn_date')}\n* **Remarks**: {updated_record.get('remarks')}"
                }
                teams.send_teams_notification(webhook_url, "update", card_data)
                
            return jsonify({"status": "success", "message": "Document details updated successfully."})
        else:
            return jsonify({"status": "error", "message": "Failed to update document details."}), 500
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/invoices/<company>/<doc_type>/<path:doc_no>', methods=['DELETE'])
@requires_auth
def delete_document(company, doc_type, doc_no):
    try:
        success = db.delete_document(company, doc_type, doc_no)
        if success:
            return jsonify({"status": "success", "message": "Document deleted successfully."})
        else:
            return jsonify({"status": "error", "message": "Document not found or could not be deleted."}), 404
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- REPORTS & DOWNLOADS ENDPOINTS ---

@app.route('/api/reports', methods=['GET'])
@requires_auth
def get_reports():
    company = request.args.get('company')
    doc_type = request.args.get('doc_type')
    
    try:
        documents = db.get_documents(company, doc_type)
        if not documents:
            return jsonify({
                "total_count": 0,
                "total_amt": 0.0,
                "pending_mrn": 0,
                "pending_po": 0,
                "pending_eway": 0,
                "done_mrn": 0,
                "by_company": [],
                "by_date": []
            })
            
        def calculate_doc_volume(doc):
            amt = doc.get('amount') or 0.0
            if (not amt or amt == 0.0) and doc.get('doc_type') == 'Challan':
                qtys = str(doc.get('deliver_qty') or '').split('\n')
                prices = str(doc.get('unit_price') or '').split('\n')
                chal_amt = 0.0
                for q_str, p_str in zip(qtys, prices):
                    try:
                        q = float(q_str.replace(',', '').strip())
                        p = float(p_str.replace(',', '').strip())
                        chal_amt += q * p
                    except (ValueError, TypeError):
                        pass
                return chal_amt
            return float(amt or 0.0)

        total_count = len(documents)
        total_amt = sum(calculate_doc_volume(doc) for doc in documents)
        
        pending_po = sum(
            1 for doc in documents 
            if doc.get('doc_type') == 'Challan' 
            and (not doc.get('po_status') or any(s.strip().upper() != 'CLOSED' for s in str(doc.get('po_status')).split('\n')))
        )
        pending_mrn = sum(1 for doc in documents if str(doc.get('mrn_status', '')).upper() != 'DONE')
        pending_eway = sum(
            1 for doc in documents 
            if doc.get('doc_type') == 'Invoice' 
            and str(doc.get('eway_bill', '')).upper() not in ['DONE', 'N/A']
            and not str(doc.get('eway_bill_no') or '').strip()
            and (doc.get('amount') or 0.0) >= 50000 
        )
        
        done_mrn = total_count - pending_mrn
        
        customers = {}
        for doc in documents:
            cust = doc.get('bill_to') or 'Unknown'
            amt = doc.get('amount') or 0.0
            customers[cust] = customers.get(cust, 0.0) + amt
            
        by_company = [{"company": k, "amount": v} for k, v in sorted(customers.items(), key=lambda x: x[1], reverse=True)[:5]]
        
        dates = {}
        for doc in documents:
            dt = doc.get('doc_date') or 'Unknown Date'
            amt = doc.get('amount') or 0.0
            dates[dt] = dates.get(dt, 0.0) + amt
            
        by_date = [{"date": k, "amount": v} for k, v in sorted(dates.items())]
        
        return jsonify({
            "total_count": total_count,
            "total_amt": total_amt,
            "pending_mrn": pending_mrn,
            "pending_po": pending_po,
            "pending_eway": pending_eway,
            "done_mrn": done_mrn,
            "by_company": by_company,
            "by_date": by_date
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/download-excel/<company>/<file_type>', methods=['GET'])
@requires_auth
def download_excel(company, file_type):
    if company not in db.COMPANIES:
        return jsonify({"status": "error", "message": "Invalid company name."}), 400
    if file_type not in ['main', 'pending']:
        return jsonify({"status": "error", "message": "Invalid file type."}), 400
        
    excel_path = db.get_company_excel_path(company, file_type)
    if not os.path.exists(excel_path):
        db.sync_to_excel(company)
        
    if not os.path.exists(excel_path):
        return jsonify({"status": "error", "message": "Excel file is empty or not created yet."}), 404
        
    file_suffix = "_Pending_MRN" if file_type == 'pending' else ""
    file_name = f"{company.replace(' ', '_')}{file_suffix}.xlsx"
    
    response = send_file(
        excel_path,
        as_attachment=True,
        download_name=file_name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response.headers["Content-Disposition"] = f'attachment; filename="{file_name}"'
    return response

@app.route('/api/download-report/<company>/<report_type>', methods=['GET'])
@requires_auth
def download_report(company, report_type):
    if company not in db.COMPANIES:
        return jsonify({"status": "error", "message": "Invalid company name."}), 400
        
    doc_type = request.args.get('doc_type', 'Invoice')
    documents = db.get_documents(company, doc_type)
    
    # Custom high-fidelity openpyxl paths for pending_mrn and pending_po
    if report_type in ["pending_mrn", "pending_po"]:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        
        wb = openpyxl.Workbook()
        ws = wb.active
        
        if report_type == "pending_mrn":
            sheet_name = "Pending MRNs"
            ws.title = sheet_name
            
            # Filter by the requested doc_type
            pending_list = [d for d in db.get_documents(company, doc_type) if str(d.get('mrn_status', '')).upper() != 'DONE']
            
            if doc_type == 'Invoice':
                cols = [
                    'S.No', 'Invoice No', 'Invoice Date', 'Challan No', 'Bill To',
                    'PO/WO No', 'Item Description', 'Invoice Amt', 'Eway Bill', 'E way Bill no', 'ASN No', 'Mode of Transp',
                    'Dock Code', 'Delivery Date', 'Delivery Site', 'MRN Status', 'MRN Date', 'MRN No', 'Remarks'
                ]
            else:
                cols = [
                    'S.No', 'Delivery Challan', 'Item Description', 'Deliver Qty', 'UOM',
                    'Unit Price', 'Delivery Date', 'Delivery Location', 'PO Status', 'PO Number',
                    'PO Qty', 'Difference Between Delivery & PO Qty', 'Invoice Number', 'Remarks'
                ]
                
            ws.append(cols)
            
            # Formatting headers
            header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
            for col_idx in range(1, len(cols) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
                
            sorted_pending = sorted(pending_list, key=lambda x: x['id'])
            
            row_idx_counter = 1
            for doc in sorted_pending:
                if doc_type == 'Invoice':
                    row_data = []
                    for col in cols:
                        if col == 'S.No':
                            row_data.append(row_idx_counter)
                        elif col == 'Invoice No':
                            row_data.append(doc.get('doc_no', ''))
                        elif col == 'Invoice Date':
                            row_data.append(doc.get('doc_date', ''))
                        elif col == 'Challan No':
                            row_data.append(doc.get('challan_no', ''))
                        elif col == 'Bill To':
                            row_data.append(doc.get('bill_to', ''))
                        elif col == 'PO/WO No':
                            row_data.append(doc.get('po_wo_no', ''))
                        elif col == 'Item Description':
                            row_data.append(doc.get('item_description', ''))
                        elif col == 'Invoice Amt':
                            try:
                                val = float(doc.get('amount') or 0.0)
                            except ValueError:
                                val = doc.get('amount')
                            row_data.append(val)
                        elif col == 'Eway Bill':
                            row_data.append(doc.get('eway_bill', ''))
                        elif col == 'E way Bill no':
                            row_data.append(doc.get('eway_bill_no', ''))
                        elif col == 'ASN No':
                            row_data.append(doc.get('asn_no', ''))
                        elif col == 'Mode of Transp':
                            row_data.append(doc.get('mode_of_transp', ''))
                        elif col == 'Dock Code':
                            row_data.append(doc.get('dock_code', ''))
                        elif col == 'Delivery Date':
                            row_data.append(doc.get('delivery_date', ''))
                        elif col == 'Delivery Site':
                            row_data.append(doc.get('delivery_site', ''))
                        elif col == 'MRN Status':
                            row_data.append(doc.get('mrn_status', 'PENDING'))
                        elif col == 'MRN Date':
                            row_data.append(doc.get('mrn_date', ''))
                        elif col == 'MRN No':
                            row_data.append(doc.get('mrn_no', ''))
                        elif col == 'Remarks':
                            row_data.append(doc.get('remarks', ''))
                        else:
                            row_data.append('')
                    ws.append(row_data)
                    row_idx_counter += 1
                else: # Delivery Challan
                    items = str(doc.get('item_description') or '').split('\n')
                    qtys = str(doc.get('deliver_qty') or '').split('\n')
                    uoms = str(doc.get('uom') or '').split('\n')
                    prices = str(doc.get('unit_price') or '').split('\n')
                    po_qtys = str(doc.get('po_qty') or '').split('\n')
                    
                    max_len = max(len(items), len(qtys), len(uoms), len(prices), len(po_qtys))
                    max_len = max(max_len, 1)
                    
                    items += [''] * (max_len - len(items))
                    qtys += [''] * (max_len - len(qtys))
                    uoms += [''] * (max_len - len(uoms))
                    prices += [''] * (max_len - len(prices))
                    po_qtys += [''] * (max_len - len(po_qtys))
                    
                    for i in range(max_len):
                        row_data = []
                        for col in cols:
                            if col == 'S.No':
                                row_data.append(row_idx_counter)
                            elif col == 'Delivery Challan':
                                row_data.append(doc.get('doc_no', ''))
                            elif col == 'Item Description':
                                row_data.append(items[i])
                            elif col == 'Deliver Qty':
                                row_data.append(qtys[i])
                            elif col == 'UOM':
                                row_data.append(uoms[i])
                            elif col == 'Unit Price':
                                row_data.append(prices[i])
                            elif col == 'Delivery Date':
                                row_data.append(doc.get('doc_date', ''))
                            elif col == 'Delivery Location':
                                row_data.append(doc.get('delivery_site', ''))
                            elif col == 'PO Status':
                                row_data.append(doc.get('po_status', ''))
                            elif col == 'PO Number':
                                row_data.append(doc.get('po_wo_no', ''))
                            elif col == 'PO Qty':
                                row_data.append(po_qtys[i])
                            elif col == 'Difference Between Delivery & PO Qty':
                                d_val = 0.0
                                p_val = 0.0
                                try:
                                    if qtys[i]:
                                        d_val = float(qtys[i].replace(',', '').strip())
                                except ValueError:
                                    pass
                                try:
                                    if po_qtys[i]:
                                        p_val = float(po_qtys[i].replace(',', '').strip())
                                except ValueError:
                                    pass
                                diff = d_val - p_val
                                diff_str = str(int(diff)) if diff.is_integer() else f"{diff:.2f}"
                                row_data.append(diff_str)
                            elif col == 'Invoice Number':
                                row_data.append(doc.get('invoice_no', ''))
                            elif col == 'Remarks':
                                row_data.append(doc.get('remarks', ''))
                            else:
                                row_data.append('')
                        ws.append(row_data)
                        row_idx_counter += 1
                        
        elif report_type == "pending_po":
            sheet_name = "Pending POs"
            ws.title = sheet_name
            
            # Filter by Challan
            pending_po_list = [d for d in db.get_documents(company, 'Challan') if str(d.get('po_status', '')).upper() != 'CLOSED']
            
            cols = [
                'S.No', 'Delivery Challan', 'Item Description', 'Deliver Qty', 'UOM',
                'Unit Price', 'Delivery Date', 'Delivery Location', 'PO Status', 'PO Number',
                'PO Qty', 'Difference Between Delivery & PO Qty', 'Invoice Number', 'Remarks'
            ]
            
            ws.append(cols)
            
            # Formatting headers
            header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
            for col_idx in range(1, len(cols) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
                
            sorted_pending = sorted(pending_po_list, key=lambda x: x['id'])
            
            row_idx_counter = 1
            for doc in sorted_pending:
                items = str(doc.get('item_description') or '').split('\n')
                qtys = str(doc.get('deliver_qty') or '').split('\n')
                uoms = str(doc.get('uom') or '').split('\n')
                prices = str(doc.get('unit_price') or '').split('\n')
                po_qtys = str(doc.get('po_qty') or '').split('\n')
                
                max_len = max(len(items), len(qtys), len(uoms), len(prices), len(po_qtys))
                max_len = max(max_len, 1)
                
                items += [''] * (max_len - len(items))
                qtys += [''] * (max_len - len(qtys))
                uoms += [''] * (max_len - len(uoms))
                prices += [''] * (max_len - len(prices))
                po_qtys += [''] * (max_len - len(po_qtys))
                
                for i in range(max_len):
                    row_data = []
                    for col in cols:
                        if col == 'S.No':
                            row_data.append(row_idx_counter)
                        elif col == 'Delivery Challan':
                            row_data.append(doc.get('doc_no', ''))
                        elif col == 'Item Description':
                            row_data.append(items[i])
                        elif col == 'Deliver Qty':
                            row_data.append(qtys[i])
                        elif col == 'UOM':
                            row_data.append(uoms[i])
                        elif col == 'Unit Price':
                            row_data.append(prices[i])
                        elif col == 'Delivery Date':
                            row_data.append(doc.get('doc_date', ''))
                        elif col == 'Delivery Location':
                            row_data.append(doc.get('delivery_site', ''))
                        elif col == 'PO Status':
                            row_data.append(doc.get('po_status', ''))
                        elif col == 'PO Number':
                            row_data.append(doc.get('po_wo_no', ''))
                        elif col == 'PO Qty':
                            row_data.append(po_qtys[i])
                        elif col == 'Difference Between Delivery & PO Qty':
                            d_val = 0.0
                            p_val = 0.0
                            try:
                                if qtys[i]:
                                    d_val = float(qtys[i].replace(',', '').strip())
                            except ValueError:
                                pass
                            try:
                                if po_qtys[i]:
                                    p_val = float(po_qtys[i].replace(',', '').strip())
                            except ValueError:
                                pass
                            diff = d_val - p_val
                            diff_str = str(int(diff)) if diff.is_integer() else f"{diff:.2f}"
                            row_data.append(diff_str)
                        elif col == 'Invoice Number':
                            row_data.append(doc.get('invoice_no', ''))
                        elif col == 'Remarks':
                            row_data.append(doc.get('remarks', ''))
                        else:
                            row_data.append('')
                    ws.append(row_data)
                    row_idx_counter += 1

        # Format rows & wrap newlines
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                if cell.column == 8 and report_type == "pending_mrn" and doc_type == 'Invoice':
                    cell.number_format = '₹#,##0.00'
                if cell.value and isinstance(cell.value, str) and '\n' in cell.value:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')

        # Auto fit column widths
        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.value is not None:
                    lines = str(cell.value).split('\n')
                    for line in lines:
                        if len(line) > max_len:
                            max_len = len(line)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 10), 100)
            
        output = io.BytesIO()
        wb.save(output)
        wb.close()
        output.seek(0)
        
        file_suffix = "pending_mrn" if report_type == "pending_mrn" else "pending_po"
        file_name = f"{company.replace(' ', '_')}_{file_suffix}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        response = send_file(
            output,
            as_attachment=True,
            download_name=file_name,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response.headers["Content-Disposition"] = f'attachment; filename="{file_name}"'
        return response

    # Standard pandas-based export for volume, sequence_gap, pending_eway
    df = pd.DataFrame()
    sheet_name = "Report"
    
    if report_type == "volume":
        if doc_type == 'Invoice':
            cols = {
                'doc_no': 'Invoice No', 'doc_date': 'Invoice Date', 'challan_no': 'Challan No',
                'bill_to': 'Bill To', 'po_wo_no': 'PO/WO No', 'amount': 'Invoice Amt', 'remarks': 'Remarks'
            }
        else:
            cols = {
                'doc_no': 'Delivery Challan No', 'item_description': 'Item Description',
                'deliver_qty': 'Deliver Qty', 'uom': 'UOM', 'unit_price': 'Unit Price',
                'delivery_date': 'Delivery Date', 'po_wo_no': 'PO Number', 'po_qty': 'PO Qty',
                'amount': 'Amount', 'remarks': 'Remarks'
            }
            
        if documents:
            df = pd.DataFrame(documents)[list(cols.keys())].rename(columns=cols)
            amt_col = 'Invoice Amt' if doc_type == 'Invoice' else 'Amount'
            df[amt_col] = pd.to_numeric(df[amt_col], errors='coerce').fillna(0.0)
            total_sum = df[amt_col].sum()
            
            summary = {c: "" for c in df.columns}
            summary[df.columns[0]] = "TOTAL SUMMARY"
            summary[amt_col] = total_sum
            df = pd.concat([df, pd.DataFrame([summary])], ignore_index=True)
            
        sheet_name = "Volume Report"
        
    elif report_type == "sequence_gap":
        seq_check = db.check_invoice_sequence(company, doc_type)
        gaps = seq_check.get("missing_numbers", [])
        df = pd.DataFrame(gaps, columns=[f"Missing Sequential {doc_type} Numbers"])
        sheet_name = "Sequence Gaps"
        
    elif report_type == "pending_eway":
        pending_eway_list = [
            d for d in documents 
            if d.get('doc_type') == 'Invoice' 
            and (d.get('amount') or 0.0) >= 50000 
            and not str(d.get('eway_bill_no') or '').strip()
        ]
        if pending_eway_list:
            df = pd.DataFrame(pending_eway_list)[[
                'doc_no', 'doc_date', 'bill_to', 'po_wo_no', 'item_description', 'amount', 'eway_bill_no', 'remarks'
            ]].rename(columns={
                'doc_no': 'Invoice No',
                'doc_date': 'Invoice Date',
                'bill_to': 'Bill To',
                'po_wo_no': 'PO/WO No',
                'item_description': 'Item Description',
                'amount': 'Invoice Amt',
                'eway_bill_no': 'E way Bill no',
                'remarks': 'Remarks'
            })
        sheet_name = "Pending Eway Bills"
        
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
    output.seek(0)
    
    file_name = f"{company.replace(' ', '_')}_{report_type}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response = send_file(
        output,
        as_attachment=True,
        download_name=file_name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response.headers["Content-Disposition"] = f'attachment; filename="{file_name}"'
    return response


def is_local_request(remote_addr):
    import socket
    local_ips = {'127.0.0.1', 'localhost', '::1'}
    try:
        hostname = socket.gethostname()
        local_ips.add(socket.gethostbyname(hostname))
        for ip_info in socket.getaddrinfo(hostname, None):
            local_ips.add(ip_info[4][0])
    except Exception:
        pass
    return remote_addr in local_ips

@app.route('/api/open-excel/<company>/<file_type>', methods=['POST'])
@requires_auth
def open_excel(company, file_type):
    if company not in db.COMPANIES:
        return jsonify({"status": "error", "message": "Invalid company name."}), 400
    if file_type not in ['main', 'pending']:
        return jsonify({"status": "error", "message": "Invalid file type."}), 400
        
    excel_path = db.get_company_excel_path(company, file_type)
    if not os.path.exists(excel_path):
        db.sync_to_excel(company)
        
    if not os.path.exists(excel_path):
        return jsonify({"status": "error", "message": "Excel file could not be created."}), 404
        
    # Detect remote network access vs local access
    is_local = is_local_request(request.remote_addr)
    
    client_host = request.args.get('client_host', '').split(':')[0].strip().lower()
    request_host = request.headers.get('Host', '').split(':')[0].strip().lower()
    local_hostnames = {'localhost', '127.0.0.1', 'invoice.local', 'muskan', '::1'}
    
    if client_host and client_host not in local_hostnames:
        is_local = False
    if request_host and request_host not in local_hostnames:
        is_local = False
        
    token = request.headers.get("Authorization", "").split(" ")[1] if request.headers.get("Authorization") else (request.args.get("token") or "")
    
    if not is_local:
        return jsonify({
            "status": "download",
            "message": "Remote connection detected. Downloading file to your device...",
            "url": f"/api/download-excel/{company}/{file_type}?token={token}"
        })
        
    try:
        label = "Pending MRNs" if file_type == 'pending' else "Main Ledger"
        if hasattr(os, 'startfile'):
            os.startfile(excel_path)
            return jsonify({"status": "success", "message": f"Opening {company} {label} in Microsoft Excel..."})
        else:
            subprocess.run(['xdg-open', excel_path], check=True)
            return jsonify({"status": "success", "message": f"Opening {company} {label}..."})
    except Exception as e:
        # Fallback to download if starting application failed (e.g. no Excel installed)
        return jsonify({
            "status": "download",
            "message": "Could not open Excel locally. Downloading file instead...",
            "url": f"/api/download-excel/{company}/{file_type}?token={token}"
        })

def register_mdns_service():
    try:
        import socket
        from zeroconf import Zeroconf, ServiceInfo
        import atexit

        # Find local IP address
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        local_ip = s.getsockname()[0]
        s.close()

        zeroconf = Zeroconf()
        service_type = "_http._tcp.local."
        service_name = "InvoSync Server._http._tcp.local."

        info = ServiceInfo(
            type_=service_type,
            name=service_name,
            addresses=[socket.inet_aton(local_ip)],
            port=5000,
            properties={},
            server="invoice.local."
        )
        zeroconf.register_service(info)
        print(f"Registered mDNS service 'invoice.local' pointing to {local_ip}:5000")
        
        # Clean up on exit
        atexit.register(lambda: (zeroconf.unregister_service(info), zeroconf.close()))
    except Exception as e:
        print(f"Failed to register mDNS service: {e}")

if __name__ == '__main__':
    # Register mDNS service so http://invoice.local:5000/ resolves on the local network
    if os.environ.get('WERKZEUG_RUN_MAIN') == 'true' or 'WERKZEUG_RUN_MAIN' not in os.environ:
        register_mdns_service()

    # Bind to 0.0.0.0 for local network access (reliable IPv4)
    app.run(host='0.0.0.0', port=5000, debug=True)
